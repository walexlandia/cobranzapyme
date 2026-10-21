"""Checks that run without network access or a PostgreSQL driver."""

import ast
import hashlib
import importlib.util
import io
import json
import os
import sys
import types
import uuid
from datetime import date, datetime, timezone
from decimal import Decimal
from pathlib import Path
from zoneinfo import ZoneInfo
from zipfile import ZipFile

from openpyxl import Workbook
from pydantic import BaseModel, ValidationError
from sqlalchemy.dialects import postgresql
from sqlalchemy.schema import CreateTable


ROOT = Path(__file__).resolve().parents[1]
sys.path.insert(0, str(ROOT / "backend"))

from app.base import Base  # noqa: E402
from app import models  # noqa: E402,F401
from app.config_validation import real_environment_issues  # noqa: E402
from app.domain import InvoiceStatus  # noqa: E402
from app.schemas import (  # noqa: E402
    ContractRules,
    ClassificationPolicyUpdate,
    ExceptionReviewCommand,
    ExceptionResponseDraftCreate,
    DeliveryStatusUpdate,
    InboundMessageCreate,
    ManualResponseCreate,
    PayerProfileProposalChanges,
    PayerProfileView,
    PayerRequirements,
    PreventiveConfirmationCreate,
    MessageTemplateCreate,
)
from app.services.classification import (  # noqa: E402
    _classification_result,
    _extract_amounts,
    _extract_dates,
)
from app.services.communications import safety_flags, validate_template_text  # noqa: E402
from app.services.client_portal import build_client_portal_summary  # noqa: E402
from app.services.invoice_import import (  # noqa: E402
    parse_contact_file,
    parse_invoice_file,
    parse_sii_status_file,
)
from app.services.investor_package import build_investor_package  # noqa: E402
from app.services.final_report import (  # noqa: E402
    amount_band,
    anonymized_alias,
    build_economics_scenarios,
    evaluate_g4,
    prioritize_backlog,
)
from app.services.payability import checklist_status, requirement_specs  # noqa: E402
from app.services.payments import parse_bank_statement_file  # noqa: E402
from app.services.pilot_metrics import (  # noqa: E402
    _mode_breakdown,
    _percent,
    build_pilot_metrics,
)
from app.services.operational_health import (  # noqa: E402
    OperationalMetrics,
    evaluate_operational_health,
    worker_delay_seconds,
)
from app.services.workflow_engine import determine_stage, next_allowed_at, stage_windows  # noqa: E402
from app.utils.rut import normalize_rut  # noqa: E402


def check_rut() -> None:
    assert normalize_rut("76.086.428-5") == "76086428-5"
    try:
        normalize_rut("76.086.428-4")
    except ValueError:
        pass
    else:
        raise AssertionError("An invalid RUT was accepted")


def check_csv_import() -> None:
    content = (
        "RUT receptor;Razón social;Folio;Fecha emisión;Fecha vencimiento;Monto total;OC\n"
        "76.086.428-5;Pagador Demo;F-100;01-10-2026;31-10-2026;$1.250.000;OC-88\n"
    ).encode()
    rows, issues = parse_invoice_file("facturas.csv", content)
    assert not issues, issues
    assert rows[0].total_amount == 1_250_000
    assert rows[0].payer_rut == "76086428-5"


def check_xlsx_import() -> None:
    workbook = Workbook()
    sheet = workbook.active
    sheet.append(["RUT receptor", "Razón social", "Folio", "Fecha emisión", "Fecha vencimiento", "Monto total"])
    sheet.append(["76.086.428-5", "Pagador Demo", "F-200", "2026-10-01", "2026-10-31", 900000])
    payload = io.BytesIO()
    workbook.save(payload)
    rows, issues = parse_invoice_file("facturas.xlsx", payload.getvalue())
    assert not issues, issues
    assert rows[0].folio == "F-200"


def check_contact_import() -> None:
    content = (
        "RUT pagador;Razón social;Nombre contacto;Rol;Correo;Teléfono\n"
        "76.086.428-5;Pagador Demo;María Demo;Cuentas por pagar;maria@example.com;+56911111111\n"
    ).encode()
    rows, issues = parse_contact_file("contactos.csv", content)
    assert not issues, issues
    assert rows[0].role == "accounts_payable"
    assert str(rows[0].email) == "maria@example.com"


def check_sii_status_import() -> None:
    content = (
        "Tipo documento;Folio;Estado SII;Fecha consulta;Plazo reclamo;Detalle\n"
        "factura;F-100;Reclamada;2026-10-14;2026-10-15;OC inválida\n"
    ).encode()
    rows, issues = parse_sii_status_file("sii.csv", content)
    assert not issues, issues
    assert rows[0].sii_status == "claimed"
    assert rows[0].claim_deadline == date(2026, 10, 15)


def check_bank_statement_import() -> None:
    content = (
        "id movimiento,fecha,monto,descripcion,rut contraparte,moneda\n"
        "TX-DEMO-001,2026-10-15,1250000,Transferencia ficticia,76.086.428-5,CLP\n"
    ).encode()
    rows, issues = parse_bank_statement_file("cartola.csv", content)
    assert not issues, issues
    assert rows[0].external_id == "TX-DEMO-001"
    assert rows[0].counterparty_rut == "76086428-5"
    assert rows[0].amount == Decimal("1250000.00")


def check_payability_rules() -> None:
    invoice = models.Invoice(
        id=uuid.uuid4(),
        client_id=uuid.uuid4(),
        payer_id=uuid.uuid4(),
        folio="PAY-1",
        issued_on=date(2026, 10, 1),
        due_on=date(2026, 10, 31),
        total_amount=100000,
        status=InvoiceStatus.ISSUED,
    )
    requirements = PayerRequirements(
        requires_purchase_order=True,
        requires_hes=True,
    )
    specs = requirement_specs(invoice, requirements)
    blocking = [spec for spec in specs if spec["required"] and spec["blocks_payability"]]
    items = [
        models.PayabilityItem(
            client_id=invoice.client_id,
            invoice_id=invoice.id,
            item_type=spec["item_type"],
            label=spec["label"],
            status="fulfilled",
            required=spec["required"],
            blocks_payability=spec["blocks_payability"],
            responsible_role=spec["responsible_role"],
            evidence={},
        )
        for spec in blocking
    ]
    status, completed, total, missing = checklist_status(
        invoice, items, requirements=requirements, today=date(2026, 10, 20)
    )
    assert status == "payable"
    assert completed == total == len(blocking)
    assert not missing


def check_sii_deadline_alert() -> None:
    invoice = models.Invoice(
        id=uuid.uuid4(),
        client_id=uuid.uuid4(),
        payer_id=uuid.uuid4(),
        folio="SII-RISK-1",
        issued_on=date(2026, 1, 10),
        due_on=date(2026, 2, 9),
        total_amount=100000,
        status=InvoiceStatus.ISSUED,
    )
    requirements = PayerRequirements(risk_days_before_due=7)
    item = models.PayabilityItem(
        client_id=invoice.client_id,
        invoice_id=invoice.id,
        item_type="sii_acceptance",
        label="Aceptación en SII",
        status="pending",
        required=True,
        blocks_payability=True,
        responsible_role="client",
        evidence={"sii_status": "pending", "claim_deadline": "2026-01-20"},
    )
    status, _, _, _ = checklist_status(
        invoice,
        [item],
        requirements=requirements,
        today=date(2026, 1, 25),
    )
    assert status == "at_risk"


def check_preventive_confirmation_contract() -> None:
    valid = PreventiveConfirmationCreate(
        approval_status="approved",
        approved_at=datetime(2026, 1, 25, 13, 0, tzinfo=timezone.utc),
        scheduled_payment_date=date(2026, 2, 5),
        reference="Confirmación CxP 1842",
    )
    assert valid.scheduled_payment_date == date(2026, 2, 5)
    invalid_payloads = [
        {
            "approval_status": "approved",
            "approved_at": datetime(2026, 1, 25, 13, 0, tzinfo=timezone.utc),
        },
        {
            "approval_status": "rejected",
            "approved_at": datetime(2026, 1, 25, 13, 0, tzinfo=timezone.utc),
            "scheduled_payment_date": date(2026, 2, 5),
        },
    ]
    for payload in invalid_payloads:
        try:
            PreventiveConfirmationCreate(**payload)
        except ValidationError:
            continue
        raise AssertionError(f"An invalid preventive confirmation was accepted: {payload}")


def check_global_payer_profile_contract() -> None:
    public_fields = set(PayerProfileView.model_fields)
    assert public_fields == {
        "id",
        "payer_id",
        "version",
        "updated_at",
        "operational_requirements",
        "supplier_portal_url",
        "supplier_portal_instructions",
        "corporate_channels",
        "payment_cycle",
    }
    forbidden = {
        "client_id",
        "invoice_id",
        "amount",
        "private_notes",
        "personal_contacts",
        "conversations",
    }
    assert public_fields.isdisjoint(forbidden)


def check_communication_safety_contract() -> None:
    validate_template_text(
        "Seguimiento · {client_name}",
        "Hola {recipient_name}, revisamos {invoice_summary}. {action_request}",
    )
    assert safety_flags(None, "Iniciaremos una demanda y embargo") == [
        "threat",
        "legal_pressure",
    ]
    try:
        validate_template_text(None, "Hola {unknown_field}, mensaje suficientemente largo")
    except ValueError:
        pass
    else:
        raise AssertionError("An unknown communication template field was accepted")
    payload = MessageTemplateCreate(
        client_id=uuid.uuid4(),
        template_key="payment_due",
        channel="email",
        stage="due",
        recipient_role="accounts_payable",
        tone="cordial",
        subject_template="Factura · {payer_name}",
        body_template="Hola {recipient_name}, {invoice_summary}",
    )
    assert payload.approved is False


def check_inbound_communication_contract() -> None:
    client_id = uuid.uuid4()
    invoice_id = uuid.uuid4()
    inbound = InboundMessageCreate(
        client_id=client_id,
        channel="email",
        provider_id="local-webhook-001",
        sender="cuentas.por.pagar@example.test",
        body="Factura recibida y programada para pago.",
        invoice_folios=["DEMO-001"],
        attachments=[{"original_name": "acuse.pdf"}],
    )
    manual = ManualResponseCreate(
        client_id=client_id,
        invoice_ids=[invoice_id],
        channel="phone",
        body="Confirmación registrada por teléfono.",
    )
    status = DeliveryStatusUpdate(client_id=client_id, status="read")
    assert inbound.attachments[0].kind == "reply_attachment"
    assert manual.invoice_ids == [invoice_id]
    assert status.status == "read"


def check_local_classification_rules() -> None:
    category, confidence, evidence = _classification_result(
        "Confirmamos que el pago quedó programado para el 24/10/2026 por $1.250.000."
    )
    assert category == "payment_scheduled"
    assert confidence >= 0.9
    assert evidence["ambiguous"] is False
    assert _extract_dates("Pagaremos el 24/10/2026 y no el 2026-10-25") == [
        "2026-10-24",
        "2026-10-25",
    ]
    assert _extract_amounts("Monto CLP 1.250.000") == [1250000]
    category, confidence, _ = _classification_result("Gracias por la información.")
    assert category == "other"
    assert confidence < 0.9
    assert _classification_result("El pago fue efectuado esta mañana.")[0] == "paid"
    assert _classification_result("La orden de compra está pendiente.")[0] == "missing_oc_hes"
    assert _classification_result("Los datos bancarios están incorrectos.")[0] == "supplier_data_issue"


def check_supervision_contract() -> None:
    client_id = uuid.uuid4()
    review = ExceptionReviewCommand(
        client_id=client_id,
        category="payment_scheduled",
        extracted_data={"payment_dates": ["2026-10-24"]},
        human_minutes="4.5",
    )
    assert review.checklist_updates == []
    assert review.workflow_stage is None
    assert ClassificationPolicyUpdate(
        client_id=client_id, mode="supervised"
    ).mode == "supervised"
    assert PayerProfileProposalChanges.model_validate(
        {"payment_cycle": "Todos los viernes"}
    ).model_dump(exclude_unset=True) == {"payment_cycle": "Todos los viernes"}
    try:
        PayerProfileProposalChanges.model_validate({"campo_inventado": "no permitido"})
    except ValidationError:
        pass
    else:
        raise AssertionError("Profile proposals must reject unknown fields")
    response = ExceptionResponseDraftCreate(
        client_id=client_id,
        channel="email",
        subject="Respuesta segura",
        body="Contenido preparado para revisión humana.",
    )
    assert response.channel == "email"


def check_openapi_contract() -> None:
    """Compile the HTTP contract without loading the incompatible global settings package."""

    shim = types.ModuleType("pydantic_settings")

    class OfflineBaseSettings(BaseModel):
        def __init__(self, **data: object) -> None:
            for name in type(self).model_fields:
                value = os.getenv(name.upper())
                if value is not None:
                    data.setdefault(name, value)
            super().__init__(**data)

    shim.BaseSettings = OfflineBaseSettings
    shim.SettingsConfigDict = dict
    previous_module = sys.modules.get("pydantic_settings")
    previous_url = os.environ.get("DATABASE_URL")
    try:
        sys.modules["pydantic_settings"] = shim
        os.environ["DATABASE_URL"] = "sqlite+pysqlite:///:memory:"
        from app.main import app

        schema = app.openapi()
        assert len(schema["paths"]) >= 38
        assert (
            "/api/v1/payability/invoices/{invoice_id}/preventive-confirmation"
            in schema["paths"]
        )
        assert "PreventiveConfirmationCreate" in schema["components"]["schemas"]
        assert "/api/v1/communications/messages" in schema["paths"]
        assert "/api/v1/communications/templates" in schema["paths"]
        assert "/api/v1/communications/inbound/simulate" in schema["paths"]
        assert "/api/v1/communications/manual-responses" in schema["paths"]
        assert (
            "/api/v1/communications/messages/{message_id}/delivery-status"
            in schema["paths"]
        )
        assert "/api/v1/exceptions" in schema["paths"]
        assert "/api/v1/exceptions/{exception_id}/review" in schema["paths"]
        assert "/api/v1/evaluation/sets" in schema["paths"]
        assert "/api/v1/evaluation/report" in schema["paths"]
        assert "/api/v1/evaluation/report.csv" in schema["paths"]
        assert "/api/v1/classification-policies" in schema["paths"]
        assert "/api/v1/payer-profile-proposals" in schema["paths"]
        assert (
            "/api/v1/communications/exception-responses/{exception_id}/draft"
            in schema["paths"]
        )
        assert "/api/v1/payments/manual" in schema["paths"]
        assert "/api/v1/payments/bank-statements" in schema["paths"]
        assert "/api/v1/payments/reconciliation-proposals" in schema["paths"]
        assert "/api/v1/payments/commission-report" in schema["paths"]
        assert "/api/v1/payments/commission-report.csv" in schema["paths"]
        assert "CommissionReportView" in schema["components"]["schemas"]
        assert "/api/v1/access-logs" in schema["paths"]
        assert "/api/v1/pilot/authorizations" in schema["paths"]
        assert "/api/v1/pilot/readiness" in schema["paths"]
        assert "/api/v1/client-portal/summary" in schema["paths"]
        assert "/api/v1/client-portal/weekly-report" in schema["paths"]
        assert "/api/v1/client-portal/weekly-report.csv" in schema["paths"]
        assert "/api/v1/client-portal/documents" in schema["paths"]
        assert "/api/v1/client-portal/document-submissions" in schema["paths"]
        assert (
            "/api/v1/client-portal/document-submissions/{submission_id}/review"
            in schema["paths"]
        )
        assert "ClientPortalSummaryView" in schema["components"]["schemas"]
        assert "WeeklyClientReportView" in schema["components"]["schemas"]
        assert "ClientDocumentSubmissionView" in schema["components"]["schemas"]
        assert "/api/v1/metrics/pilot" in schema["paths"]
        assert "/api/v1/metrics/pilot.csv" in schema["paths"]
        assert "/api/v1/metrics/human-effort" in schema["paths"]
        assert "/api/v1/metrics/human-effort.csv" in schema["paths"]
        assert "/api/v1/metrics/human-effort.md" in schema["paths"]
        assert "/api/v1/metrics/client-trend" in schema["paths"]
        assert "PilotMetricsView" in schema["components"]["schemas"]
        assert "HumanEffortDiagnosticView" in schema["components"]["schemas"]
        assert "/api/v1/pilot/log" in schema["paths"]
        assert "/api/v1/pilot/log/{entry_id}/resolve" in schema["paths"]
        assert "/api/v1/pilot/demo-walkthrough" in schema["paths"]
        assert "/api/v1/client-portal/escalated-exceptions" in schema["paths"]
        assert (
            "/api/v1/client-portal/escalated-exceptions/{exception_id}/response"
            in schema["paths"]
        )
        assert "PilotLogEntryView" in schema["components"]["schemas"]
        assert "/api/v1/metrics/closures" in schema["paths"]
        assert "/api/v1/metrics/closures/{closure_id}/decision" in schema["paths"]
        assert "PilotClosureSnapshotView" in schema["components"]["schemas"]
        assert "/api/v1/operations/observability" in schema["paths"]
        assert "OperationsSnapshotView" in schema["components"]["schemas"]
    finally:
        if previous_module is None:
            sys.modules.pop("pydantic_settings", None)
        else:
            sys.modules["pydantic_settings"] = previous_module
        if previous_url is None:
            os.environ.pop("DATABASE_URL", None)
        else:
            os.environ["DATABASE_URL"] = previous_url


def check_schema_compiles_for_postgresql() -> None:
    expected = {
        "clients", "users", "contracts", "payers", "payer_profiles",
        "client_payer_relations", "contacts", "invoices", "payability_items",
        "documents", "messages", "payments", "bank_transactions",
        "exception_cases", "workflow_states", "workflow_controls", "workflow_jobs", "events",
        "payer_profile_revisions", "delay_causes",
        "message_templates", "communication_threads", "message_invoices",
        "message_documents",
        "message_classifications",
        "evaluation_sets", "evaluation_examples", "classification_policies",
        "payer_profile_proposals", "reconciliation_proposals", "reconciliation_items",
        "payment_attributions",
        "data_access_logs",
        "pilot_authorization_events",
        "client_document_submissions",
        "pilot_log_entries",
        "pilot_closure_snapshots",
        "service_heartbeats",
    }
    assert set(Base.metadata.tables) == expected
    dialect = postgresql.dialect()
    for table in Base.metadata.sorted_tables:
        sql = str(CreateTable(table).compile(dialect=dialect))
        assert "CREATE TABLE" in sql
    private_tables = expected - {
        "clients", "users", "payers", "payer_profiles", "payer_profile_revisions",
        "pilot_closure_snapshots", "service_heartbeats",
    }
    for table_name in private_tables:
        assert "client_id" in Base.metadata.tables[table_name].columns


def check_migration_chain() -> None:
    migration_dir = ROOT / "backend" / "migrations" / "versions"
    revisions: dict[str, tuple[str | None, Path]] = {}
    for path in sorted(migration_dir.glob("*.py")):
        tree = ast.parse(path.read_text(encoding="utf-8"), filename=str(path))
        assignments: dict[str, object] = {}
        functions = set()
        for node in tree.body:
            if isinstance(node, (ast.Assign, ast.AnnAssign)):
                target = node.targets[0] if isinstance(node, ast.Assign) else node.target
                value = node.value
                if isinstance(target, ast.Name) and target.id in {"revision", "down_revision"}:
                    assignments[target.id] = ast.literal_eval(value)
            elif isinstance(node, ast.FunctionDef):
                functions.add(node.name)
        revision = assignments.get("revision")
        down_revision = assignments.get("down_revision")
        assert isinstance(revision, str), f"Missing revision in {path.name}"
        assert revision not in revisions, f"Duplicate migration revision {revision}"
        assert down_revision is None or isinstance(down_revision, str)
        assert {"upgrade", "downgrade"}.issubset(functions), f"Incomplete migration {path.name}"
        revisions[revision] = (down_revision, path)
    assert revisions, "No migrations found"
    heads = set(revisions)
    for down_revision, path in revisions.values():
        if down_revision is not None:
            assert down_revision in revisions, f"Unknown parent {down_revision} in {path.name}"
            heads.discard(down_revision)
    assert heads == {"0017"}, f"Unexpected migration heads: {sorted(heads)}"
    current: str | None = "0017"
    visited = []
    while current is not None:
        assert current not in visited, f"Migration cycle at {current}"
        visited.append(current)
        current = revisions[current][0]
    assert len(visited) == len(revisions), "Disconnected migration branch"


def check_migrations_match_models() -> None:
    class OperationRecorder:
        def __init__(self) -> None:
            self.tables: dict[str, set[str]] = {}

        def create_table(self, name: str, *elements: object, **_: object) -> None:
            assert name not in self.tables, f"Migration creates table twice: {name}"
            self.tables[name] = {
                element.name
                for element in elements
                if hasattr(element, "name") and element.__class__.__name__ == "Column"
            }

        def add_column(self, table: str, column: object) -> None:
            assert table in self.tables, f"Migration adds column to unknown table: {table}"
            name = getattr(column, "name")
            assert name not in self.tables[table], f"Migration adds column twice: {table}.{name}"
            self.tables[table].add(name)

        def __getattr__(self, name: str):
            if name.startswith(("create_", "execute")):
                return lambda *args, **kwargs: None
            raise AttributeError(name)

    recorder = OperationRecorder()
    alembic_module = types.ModuleType("alembic")
    alembic_module.op = recorder
    previous = sys.modules.get("alembic")
    try:
        sys.modules["alembic"] = alembic_module
        migration_paths = sorted((ROOT / "backend" / "migrations" / "versions").glob("*.py"))
        for index, path in enumerate(migration_paths):
            spec = importlib.util.spec_from_file_location(f"offline_migration_{index}", path)
            assert spec is not None and spec.loader is not None
            module = importlib.util.module_from_spec(spec)
            spec.loader.exec_module(module)
            module.upgrade()
    finally:
        if previous is None:
            sys.modules.pop("alembic", None)
        else:
            sys.modules["alembic"] = previous
    model_tables = {
        name: set(table.columns.keys()) for name, table in Base.metadata.tables.items()
    }
    assert set(recorder.tables) == set(model_tables), (
        f"Migration/model table mismatch: migrations={sorted(recorder.tables)} "
        f"models={sorted(model_tables)}"
    )
    mismatches = {
        table: {
            "missing_in_migrations": sorted(model_tables[table] - recorder.tables[table]),
            "missing_in_models": sorted(recorder.tables[table] - model_tables[table]),
        }
        for table in model_tables
        if recorder.tables[table] != model_tables[table]
    }
    assert not mismatches, f"Migration/model column mismatch: {mismatches}"


def check_phase6_operational_assets() -> None:
    compose = (ROOT / "compose.yaml").read_text(encoding="utf-8")
    secure = (ROOT / "compose.secure.yaml").read_text(encoding="utf-8")
    nginx = (ROOT / "infra" / "nginx" / "recupera.conf").read_text(encoding="utf-8")
    assert "sleep 86400" in compose
    assert "pg_dump --format=custom" in compose
    assert "DOCUMENT_ENCRYPTION_KEY" in secure and ":?" in secure
    assert "REPORT_ANONYMIZATION_KEY" in secure
    assert "ports: !reset []" in secure
    assert "ssl_protocols TLSv1.2 TLSv1.3" in nginx
    assert (ROOT / "scripts" / "backup-postgres.ps1").is_file()
    assert (ROOT / "scripts" / "verify-backup-restore.ps1").is_file()
    assert (ROOT / "scripts" / "generate-local-tls.py").is_file()
    assert (ROOT / "docs" / "fase-6" / "checklist-g3.md").is_file()


def check_phase7_client_portal_assets() -> None:
    router = (ROOT / "backend" / "app" / "routers" / "client_portal.py").read_text(
        encoding="utf-8"
    )
    service = (ROOT / "backend" / "app" / "services" / "client_portal.py").read_text(
        encoding="utf-8"
    )
    interface = (ROOT / "web" / "app" / "page.jsx").read_text(encoding="utf-8")
    tests = (ROOT / "backend" / "tests" / "test_client_portal.py").read_text(
        encoding="utf-8"
    )
    document_service = (
        ROOT / "backend" / "app" / "services" / "client_documents.py"
    ).read_text(encoding="utf-8")
    document_tests = (
        ROOT / "backend" / "tests" / "test_client_documents.py"
    ).read_text(encoding="utf-8")
    assert "_authorized_client" in router
    assert "Payment.verified.is_(True)" in service
    assert 'item.responsible_role == "client"' in service
    assert "ClientPortalPanel" in interface
    assert "weekly-report.csv" in interface
    assert "DocumentUploadDialog" in interface and "REVISIÓN HUMANA" in interface
    assert "encrypt_document_content" in document_service
    assert "content_sha256" in document_service
    assert "portal-other" in tests and "status_code == 403" in tests
    assert "decrypt_document_content" in document_tests
    assert "client_document.reviewed" in document_tests


def check_phase7_pilot_metrics_assets() -> None:
    router = (ROOT / "backend" / "app" / "routers" / "metrics.py").read_text(
        encoding="utf-8"
    )
    service = (ROOT / "backend" / "app" / "services" / "pilot_metrics.py").read_text(
        encoding="utf-8"
    )
    interface = (ROOT / "web" / "app" / "page.jsx").read_text(encoding="utf-8")
    tests = (ROOT / "backend" / "tests" / "test_pilot_metrics.py").read_text(
        encoding="utf-8"
    )
    assert "build_pilot_metrics" in router and "utf-8-sig" in router
    assert "Payment.verified.is_(True)" in service
    assert "human_hourly_cost_clp" in service and "usd_clp_rate" in service
    assert "MetricsPanel" in interface and "metrics/client-trend" in interface
    assert "economics" in tests and "status_code == 403" in tests
    assert _percent(1, 2) == Decimal("50.00")
    assert _percent(0, 0) is None
    assert _mode_breakdown(3, 1) == {
        "automatic": 3,
        "manual": 1,
        "total": 4,
        "automatic_percent": Decimal("75.00"),
    }


def check_phase7_pilot_operations_assets() -> None:
    pilot_router = (ROOT / "backend" / "app" / "routers" / "pilot.py").read_text(
        encoding="utf-8"
    )
    portal_router = (
        ROOT / "backend" / "app" / "routers" / "client_portal.py"
    ).read_text(encoding="utf-8")
    interface = (ROOT / "web" / "app" / "page.jsx").read_text(encoding="utf-8")
    seed = (ROOT / "backend" / "app" / "seed.py").read_text(encoding="utf-8")
    tests = (
        ROOT / "backend" / "tests" / "test_pilot_operations.py"
    ).read_text(encoding="utf-8")
    assert "pilot.log_recorded" in pilot_router and "pilot.log_resolved" in pilot_router
    assert "exception.client_responded" in portal_router
    assert "PilotOperationsPanel" in interface and "recupera_demo_date" in interface
    assert "PilotLogEntry(" in seed and "exception.escalated_to_client" in seed
    assert "status_code == 403" in tests and "status_code == 404" in tests
    for name in ("runbook-operacion.md", "guion-demostracion.md", "checklist-g4.md"):
        assert (ROOT / "docs" / "fase-7" / name).is_file()


def check_phase8_final_report_assets() -> None:
    router = (ROOT / "backend" / "app" / "routers" / "metrics.py").read_text(
        encoding="utf-8"
    )
    service = (ROOT / "backend" / "app" / "services" / "final_report.py").read_text(
        encoding="utf-8"
    )
    interface = (ROOT / "web" / "app" / "page.jsx").read_text(encoding="utf-8")
    tests = (ROOT / "backend" / "tests" / "test_final_report.py").read_text(
        encoding="utf-8"
    )
    migration = (
        ROOT / "backend" / "migrations" / "versions" / "0015_pilot_closure.py"
    ).read_text(encoding="utf-8")
    assert "anonymized-dataset.csv" in router and "final-report.md" in router
    assert "closures/{closure_id}/package.zip" in router
    assert "pilot.closure_prepared" in router and "pilot.g4_decided" in router
    assert "HMAC-SHA256 aliases" in service and "minimum_group_size" in service
    assert "FinalReportPanel" in interface and "Matriz G4" in interface
    assert "Sensibilidad de margen" in interface and "Snapshots y decisiones G4" in interface
    assert "cannot_approve_synthetic_g4" in tests and "status_code == 403" in tests
    assert "protect_pilot_closure_snapshot" in migration
    assert (ROOT / "scripts" / "postgres_closure_checks.sql").is_file()
    assert amount_band(Decimal("499999")) == "menos_de_500k"
    assert amount_band(Decimal("500000")) == "500k_a_2m"
    identifier = uuid.UUID("11111111-1111-1111-1111-111111111111")
    key = "offline-check-key-with-at-least-32-characters"
    assert anonymized_alias("client", identifier, key) == anonymized_alias(
        "client", identifier, key
    )
    assert anonymized_alias("client", identifier, key).startswith("PYME-")
    hypotheses = [
        {"key": "outcome_improvement", "state": "confirmed"},
        {"key": "documentary_relevance", "state": "confirmed"},
        {"key": "positive_unit_economics", "state": "confirmed"},
    ]
    synthetic = evaluate_g4(
        evidence_kind="synthetic",
        authorization_coverage_complete=False,
        period_days=56,
        client_count=3,
        settled_invoices=30,
        hypotheses=hypotheses,
        payer_concentration_percent=Decimal("40"),
    )
    assert synthetic["ready"] is False and synthetic["recommendation"] == "pending"
    authorized = evaluate_g4(
        evidence_kind="authorized_pilot",
        authorization_coverage_complete=True,
        period_days=56,
        client_count=3,
        settled_invoices=30,
        hypotheses=hypotheses,
        payer_concentration_percent=Decimal("40"),
    )
    assert authorized["ready"] is True and authorized["recommendation"] == "scale"
    scenarios = build_economics_scenarios(
        {
            "commission_clp": Decimal("100000"),
            "total_cost_clp": Decimal("80000"),
            "human_minutes": Decimal("300"),
            "managed_invoices": 10,
        },
        {"human_hourly_cost_clp": Decimal("10000")},
    )
    assert [item["key"] for item in scenarios] == [
        "conservative",
        "observed",
        "efficient",
    ]
    assert scenarios[1]["margin_clp"] == Decimal("20000.00")
    backlog = prioritize_backlog(
        hypotheses,
        payer_concentration_percent=Decimal("70"),
        g4_recommendation="iterate",
    )
    assert any(item["key"] == "diversify_payers" for item in backlog)
    assert [item["priority"] for item in backlog] == list(range(1, len(backlog) + 1))
    package_report = {
        "schema_version": "fase-8.2-v1",
        "date_from": "2026-01-01",
        "date_to": "2026-02-28",
        "period_days": 59,
        "evidence_kind": "synthetic",
        "fingerprint_sha256": "a" * 64,
        "portfolio": {
            "client_count": 2,
            "managed_invoices": 10,
            "settled_invoices": 8,
            "verified_amount_clp": "800000.00",
            "average_dso": "32.00",
            "on_time_percent": "75.00",
            "human_minutes": "120.00",
            "margin_clp": "20000.00",
        },
        "client_comparisons": [],
        "segments": [
            {
                "dimension": "amount_band",
                "segment": "menos_de_500k",
                "invoice_count": 10,
                "nominal_amount_clp": "1000000.00",
                "verified_amount_clp": "800000.00",
                "settled_invoices": 8,
                "on_time_percent": "75.00",
                "average_settlement_days": "32.00",
            }
        ],
        "hypotheses": [
            {
                "key": "outcome_improvement",
                "state": "confirmed",
                "statement": "La operación mejora el resultado.",
                "evidence": "Muestra sintética.",
            }
        ],
        "economics_scenarios": [
            {
                "key": "observed",
                "label": "Observado",
                "commission_clp": "100000.00",
                "human_cost_clp": "60000.00",
                "technology_cost_clp": "20000.00",
                "total_cost_clp": "80000.00",
                "cost_per_managed_invoice_clp": "8000.00",
                "margin_clp": "20000.00",
                "margin_percent": "20.00",
                "human_minutes_factor": "1.00",
                "technology_cost_factor": "1.00",
            }
        ],
        "backlog_priorities": [
            {
                "priority": 1,
                "key": "validate_real_pilot",
                "action": "Validar piloto real",
                "rationale": "La evidencia actual es sintética.",
            }
        ],
        "g4": {
            "ready": False,
            "recommendation": "pending",
            "reason": "La evidencia sintética no habilita G4.",
            "readiness_checks": [{"key": "authorized_evidence", "passed": False}],
        },
        "privacy": {
            "method": "HMAC-SHA256 aliases + aggregated dimensions",
            "minimum_group_size": 2,
        },
        "suppressed_groups": [{"dimension": "payer", "groups": 1}],
        "authorization_coverage": {
            "participating_clients": 2,
            "authorized_clients": 0,
            "complete": False,
        },
    }
    package, manifest = build_investor_package(
        package_report, closure_status="prepared", decision=None
    )
    repeated_package, _ = build_investor_package(
        package_report, closure_status="prepared", decision=None
    )
    assert package == repeated_package and manifest["draft"] is True
    with ZipFile(io.BytesIO(package)) as archive:
        assert set(archive.namelist()) == {
            "resumen-ejecutivo.md",
            "informe-final-v10.md",
            "dataset-anonimizado.csv",
            "unit-economics.csv",
            "backlog-priorizado.csv",
            "anexo-evidencia.md",
            "manifest.json",
        }
        stored_manifest = json.loads(archive.read("manifest.json"))
        for entry in stored_manifest["files"]:
            assert hashlib.sha256(archive.read(entry["path"])).hexdigest() == entry[
                "sha256"
            ]
        assert "BORRADOR — EVIDENCIA SINTÉTICA" in archive.read(
            "resumen-ejecutivo.md"
        ).decode("utf-8-sig")
    for name in (
        "estado.md",
        "metodologia-informe.md",
        "matriz-g4.md",
        "backlog-posterior.md",
        "paquete-postulacion.md",
    ):
        assert (ROOT / "docs" / "fase-8" / name).is_file()


def check_phase8_release_readiness_assets() -> None:
    manifest_path = ROOT / "release" / "phase-8.4.json"
    manifest = json.loads(manifest_path.read_text(encoding="utf-8"))
    assert manifest["release"] == "0.8.4"
    assert manifest["database_migration_head"] == "0017"
    assert manifest["report_schema_version"] == "fase-8.2-v1"
    assert manifest["package_schema_version"] == "fase-8.3-v1"
    assert all(gate["status"] == "blocked" for gate in manifest["external_gates"])
    assert "/health/version" in manifest["health_endpoints"]
    for relative_path in manifest["release_artifacts"]:
        assert (ROOT / relative_path).is_file(), relative_path
    readiness_path = ROOT / "scripts" / "release_readiness.py"
    spec = importlib.util.spec_from_file_location("release_readiness", readiness_path)
    assert spec is not None and spec.loader is not None
    module = importlib.util.module_from_spec(spec)
    sys.modules[spec.name] = module
    spec.loader.exec_module(module)
    checks = module.repository_checks(manifest)
    assert checks and all(item.status == "PASS" for item in checks)
    example_environment = module.environment_checks(
        ROOT / ".env.example", expected_release=manifest["release"]
    )
    assert example_environment[0].status == "FAIL"
    assert "valores de ejemplo" in example_environment[0].detail
    rehearsal = (ROOT / "scripts" / "rehearse-release.ps1").read_text(
        encoding="utf-8"
    )
    assert "recupera-release-rehearsal" in rehearsal
    assert "pg_dump" in rehearsal and "pg_restore" in rehearsal
    assert "release_rehearsal_marker" in rehearsal and "down --volumes" in rehearsal
    health = (ROOT / "backend" / "app" / "routers" / "health.py").read_text(
        encoding="utf-8"
    )
    version = (ROOT / "backend" / "app" / "version.py").read_text(encoding="utf-8")
    assert 'router.get("/health/version")' in health
    assert 'APP_VERSION = "0.8.4"' in version
    assert 'DATABASE_MIGRATION_HEAD = "0017"' in version


def check_post_plan_hardening_assets() -> None:
    safe_values = {
        "ENVIRONMENT": "pilot",
        "DATABASE_URL": "postgresql+psycopg://app:strong-password@db:5432/recupera",
        "JWT_SECRET": "jwt-secret-with-more-than-thirty-two-characters",
        "DOCUMENT_ENCRYPTION_KEY": "document-key-with-more-than-thirty-two-chars",
        "REPORT_ANONYMIZATION_KEY": "report-key-with-more-than-thirty-two-characters",
        "DEMO_ADMIN_PASSWORD": "admin-password-not-used-anywhere-else",
        "CORS_ORIGINS": "",
    }
    assert real_environment_issues(safe_values) == []
    unsafe_values = {**safe_values, "JWT_SECRET": "local-development-secret-change-before-pilot"}
    assert "JWT_SECRET" in real_environment_issues(unsafe_values)
    repeated_values = {
        **safe_values,
        "REPORT_ANONYMIZATION_KEY": safe_values["DOCUMENT_ENCRYPTION_KEY"],
    }
    assert "INDEPENDENT_CRYPTO_KEYS" in real_environment_issues(repeated_values)
    assert "CORS_ORIGINS" in real_environment_issues(
        {**safe_values, "CORS_ORIGINS": "*"}
    )
    security_headers = (
        ROOT / "backend" / "app" / "security_headers.py"
    ).read_text(encoding="utf-8")
    access_logging = (ROOT / "backend" / "app" / "access_logging.py").read_text(
        encoding="utf-8"
    )
    assert "X-Content-Type-Options" in security_headers
    assert "Strict-Transport-Security" in security_headers
    assert 'request.url.path.startswith("/api/")' in security_headers
    assert 'getattr(request.state, "request_id"' in access_logging
    pilot_compose = (ROOT / "compose.pilot.yaml").read_text(encoding="utf-8")
    assert pilot_compose.count("ports: !reset []") == 2
    assert pilot_compose.count("volumes: !reset []") == 3
    assert "app.seed" not in pilot_compose and "--reload" not in pilot_compose
    gateway = (ROOT / "infra" / "nginx" / "recupera.conf").read_text(
        encoding="utf-8"
    )
    assert "X-Frame-Options DENY" in gateway
    assert "frame-ancestors 'none'" in gateway
    assert "proxy_set_header X-Request-ID $request_id" in gateway
    lock = json.loads((ROOT / "web" / "package-lock.json").read_text(encoding="utf-8"))
    assert lock["lockfileVersion"] == 3
    assert lock["packages"]["node_modules/next"]["version"]
    generator_path = ROOT / "scripts" / "generate_dependency_inventory.py"
    spec = importlib.util.spec_from_file_location("dependency_inventory", generator_path)
    assert spec is not None and spec.loader is not None
    module = importlib.util.module_from_spec(spec)
    sys.modules[spec.name] = module
    spec.loader.exec_module(module)
    assert (ROOT / "release" / "dependency-inventory.json").read_bytes() == (
        module.serialized_inventory()
    )
    verifier = (ROOT / "scripts" / "verify-local-release.ps1").read_text(
        encoding="utf-8"
    )
    assert "generate_dependency_inventory.py --check" in verifier
    assert "compose.pilot.yaml" in verifier
    assert '{{.Config.WorkingDir}}' in verifier
    assert '${webWorkdir}/app:ro' in verifier
    assert (ROOT / "docs" / "post-plan" / "modelo-amenazas.md").is_file()
    assert (ROOT / "docs" / "post-plan" / "estado-9.1.md").is_file()


def check_post_plan_observability_assets() -> None:
    healthy, healthy_alerts = evaluate_operational_health(
        OperationalMetrics(worker_observed=True, worker_status="healthy")
    )
    assert healthy == "healthy" and not healthy_alerts
    critical, alerts = evaluate_operational_health(
        OperationalMetrics(
            oldest_pending_job_age_seconds=1_801,
            legacy_plaintext_documents=1,
            worker_observed=True,
            worker_status="degraded",
            worker_stale_seconds=121,
            worker_consecutive_failures=3,
        )
    )
    assert critical == "critical"
    assert {item["key"] for item in alerts} == {
        "legacy_plaintext_documents",
        "oldest_pending_job",
        "worker_degraded",
    }
    assert [worker_delay_seconds(value) for value in (0, 1, 2, 8)] == [30, 60, 120, 300]
    migration = (
        ROOT / "backend" / "migrations" / "versions" / "0016_service_heartbeats.py"
    ).read_text(encoding="utf-8")
    assert 'revision: str = "0016"' in migration
    assert "service_heartbeats" in migration
    assert (ROOT / "scripts" / "simulate_worker_resilience.py").is_file()
    assert (ROOT / "docs" / "post-plan" / "runbook-observabilidad.md").is_file()
    assert (ROOT / "docs" / "post-plan" / "estado-9.2.md").is_file()


def check_post_plan_operations_ui_assets() -> None:
    page = (ROOT / "web" / "app" / "page.jsx").read_text(encoding="utf-8")
    styles = (ROOT / "web" / "app" / "globals.css").read_text(encoding="utf-8")
    config = (ROOT / "backend" / "app" / "config.py").read_text(encoding="utf-8")
    seed = (ROOT / "backend" / "app" / "seed.py").read_text(encoding="utf-8")

    assert "function OperationsPanel" in page
    assert 'request("/api/v1/operations/observability")' in page
    assert "window.setInterval(() => load(true), 30_000)" in page
    assert 'user?.role === "admin"' in page
    assert 'setActiveView("operations")' in page
    assert "Este panel no expone clientes" in page
    assert all(
        selector in styles
        for selector in (
            ".operations-hero",
            ".operations-grid",
            ".operational-alert-list",
            ".operations-privacy-note",
        )
    )
    assert 'demo_admin_email: str = "admin@example.com"' in config
    assert 'client_user_email = f"cliente{client_index + 1}@example.com"' in seed
    assert "LEGACY_DEMO_ADMIN_EMAIL" in seed
    assert (ROOT / "docs" / "post-plan" / "estado-9.3.md").is_file()


def check_post_plan_improvement_plan() -> None:
    plan = (
        ROOT / "docs" / "post-plan" / "plan-mejora-fases.md"
    ).read_text(encoding="utf-8")
    for phase in range(7):
        assert f"PM-{phase}" in plan
    for gate in range(6):
        assert f"PM-G{gate}" in plan
    assert "Regla obligatoria de admisión" in plan
    assert "Instrumentada localmente; gate provisional" in plan
    assert "Bloqueada por PM-G1" in plan
    assert "No autorizado en esta fase" in plan
    assert "bloqueada por dependencias externas" in plan
    posterior = (
        ROOT / "docs" / "fase-8" / "backlog-posterior.md"
    ).read_text(encoding="utf-8")
    assert "plan-mejora-fases.md" in posterior


def check_pm1_human_effort_assets() -> None:
    service = (ROOT / "backend" / "app" / "services" / "human_effort.py").read_text(
        encoding="utf-8"
    )
    generator = (ROOT / "scripts" / "generate_pm1_diagnostic.py").read_text(
        encoding="utf-8"
    )
    status = (ROOT / "docs" / "post-plan" / "estado-pm-1.md").read_text(
        encoding="utf-8"
    )
    assert 'REPORT_SCHEMA_VERSION = "pm-1-v1"' in service
    assert '"reconciliation_human_minutes"' in service
    assert "El tiempo de ciclo no se usa como sustituto" in service
    assert 'gate_status = "provisional_synthetic"' in service
    assert 'gate_status = "measurement_gaps"' in service
    assert 'gate_status = "authorization_incomplete"' in service
    migration = (
        ROOT
        / "backend"
        / "migrations"
        / "versions"
        / "0017_reconciliation_human_minutes.py"
    ).read_text(encoding="utf-8")
    assert 'revision: str = "0017"' in migration
    assert "human_minutes" in migration
    assert "human_effort_dataset_csv" in generator
    assert "human_effort_markdown" in generator
    assert "PM-G1" in status
    assert "provisional_synthetic" in status


def check_pilot_metrics_calculation() -> None:
    client_id = uuid.uuid4()
    payer_id = uuid.uuid4()
    first_id, second_id = uuid.uuid4(), uuid.uuid4()
    timestamp = datetime(2026, 2, 3, 12, tzinfo=timezone.utc)
    tenant = models.Client(
        id=client_id,
        name="Pyme métrica sintética",
        rut="TEST-METRICS",
        baseline_dso=Decimal("45"),
        baseline_on_time_percent=Decimal("35"),
    )
    invoices = [
        models.Invoice(
            id=first_id, client_id=client_id, payer_id=payer_id, folio="M-1",
            issued_on=date(2026, 1, 10), due_on=date(2026, 2, 9),
            total_amount=Decimal("100000"), status=InvoiceStatus.PAID,
        ),
        models.Invoice(
            id=second_id, client_id=client_id, payer_id=payer_id, folio="M-2",
            issued_on=date(2026, 1, 10), due_on=date(2026, 2, 9),
            total_amount=Decimal("100000"), status=InvoiceStatus.PAID,
        ),
    ]
    payments = [
        models.Payment(
            id=uuid.uuid4(), client_id=client_id, invoice_id=first_id,
            amount=Decimal("100000"), paid_on=date(2026, 2, 5), verified=True,
            source="manual", created_at=timestamp,
        ),
        models.Payment(
            id=uuid.uuid4(), client_id=client_id, invoice_id=second_id,
            amount=Decimal("100000"), paid_on=date(2026, 2, 20), verified=True,
            source="manual", created_at=timestamp,
        ),
    ]
    message_ids = [uuid.uuid4(), uuid.uuid4()]
    messages = [
        models.Message(
            id=message_ids[0], client_id=client_id, invoice_id=first_id,
            channel="email", direction="outbound", mode="automatic", body="Sintético",
            status="sent", created_at=timestamp,
        ),
        models.Message(
            id=message_ids[1], client_id=client_id, invoice_id=second_id,
            channel="email", direction="outbound", mode="manual", body="Sintético",
            status="sent", created_at=timestamp,
        ),
    ]
    links = [
        models.MessageInvoice(
            id=uuid.uuid4(), client_id=client_id, message_id=message_ids[0],
            invoice_id=first_id, allocated_llm_cost=Decimal("0.10"),
            allocated_messaging_cost=Decimal("0.02"), created_at=timestamp,
        ),
        models.MessageInvoice(
            id=uuid.uuid4(), client_id=client_id, message_id=message_ids[1],
            invoice_id=second_id, allocated_llm_cost=Decimal("0.05"),
            allocated_messaging_cost=Decimal("0.01"), created_at=timestamp,
        ),
    ]
    attributions = [
        models.PaymentAttribution(
            id=uuid.uuid4(), client_id=client_id, payment_id=payments[0].id,
            invoice_id=first_id, category="preventive", percentage=Decimal("1"),
            base_amount=Decimal("100000"), commission_amount=Decimal("1000"),
            rationale={},
        ),
        models.PaymentAttribution(
            id=uuid.uuid4(), client_id=client_id, payment_id=payments[1].id,
            invoice_id=second_id, category="recovery", percentage=Decimal("2"),
            base_amount=Decimal("100000"), commission_amount=Decimal("2000"),
            rationale={},
        ),
    ]
    exception = models.ExceptionCase(
        id=uuid.uuid4(), client_id=client_id, invoice_id=second_id,
        kind="payment_commitment", status="resolved", priority=1, detail={},
        human_minutes=Decimal("30"), resolution={}, created_at=timestamp,
    )
    cause = models.DelayCause(
        id=uuid.uuid4(), client_id=client_id, invoice_id=second_id,
        category="purchase_order", source="fixture", created_at=timestamp,
    )
    item = models.PayabilityItem(
        id=uuid.uuid4(), client_id=client_id, invoice_id=second_id,
        item_type="purchase_order", label="Orden de compra", status="fulfilled",
        required=True, blocks_payability=True, responsible_role="client", evidence={},
        created_at=datetime(2026, 2, 1, tzinfo=timezone.utc),
        updated_at=datetime(2026, 2, 2, tzinfo=timezone.utc),
    )

    class FakeSession:
        def __init__(self) -> None:
            self.results = iter(
                (
                    invoices,
                    payments,
                    [first_id, second_id],
                    [exception],
                    [],
                    links,
                    messages,
                    [],
                    [],
                    attributions,
                    [cause],
                    [item],
                )
            )

        def scalars(self, _statement: object):
            return iter(next(self.results))

    report = build_pilot_metrics(
        FakeSession(), client=tenant, date_from=date(2026, 2, 1),
        date_to=date(2026, 2, 28), human_hourly_cost_clp=Decimal("8000"),
        usd_clp_rate=Decimal("950"),
    )
    assert report["outcomes"]["average_dso"] == Decimal("33.50")
    assert report["outcomes"]["on_time_percent"] == Decimal("50.00")
    assert report["outcomes"]["recovered_amount"] == Decimal("100000.00")
    assert report["economics"]["margin_clp"] == Decimal("-1171.00")
    assert report["resolution_times"][0]["average_hours"] == Decimal("24.00")


def check_client_portal_calculation() -> None:
    client_id = uuid.uuid4()
    payer_id = uuid.uuid4()
    invoice_id = uuid.uuid4()
    client = models.Client(id=client_id, name="Pyme sintética", rut="TEST-PORTAL")
    payer = models.Payer(id=payer_id, rut="TEST-PAYER", legal_name="Pagador sintético")
    invoice = models.Invoice(
        id=invoice_id,
        client_id=client_id,
        payer_id=payer_id,
        payer=payer,
        folio="PORTAL-OFFLINE-001",
        issued_on=date(2026, 10, 1),
        due_on=date(2026, 10, 31),
        total_amount=Decimal("100000"),
        status=InvoiceStatus.ISSUED,
    )
    items = [
        models.PayabilityItem(
            id=uuid.uuid4(),
            client_id=client_id,
            invoice_id=invoice_id,
            item_type="sii_acceptance",
            label="Aceptación SII",
            status="fulfilled",
            required=True,
            blocks_payability=True,
            responsible_role="client",
            evidence={"sii_status": "accepted"},
        ),
        models.PayabilityItem(
            id=uuid.uuid4(),
            client_id=client_id,
            invoice_id=invoice_id,
            item_type="purchase_order",
            label="Orden de compra",
            status="requested",
            required=True,
            blocks_payability=True,
            responsible_role="client",
            evidence={},
        ),
    ]
    payment = models.Payment(
        id=uuid.uuid4(),
        client_id=client_id,
        invoice_id=invoice_id,
        amount=Decimal("25000"),
        paid_on=date(2026, 10, 20),
        verified=True,
        source="manual",
    )
    profile = models.PayerProfile(
        id=uuid.uuid4(),
        payer_id=payer_id,
        operational_requirements={
            "requires_purchase_order": True,
            "requires_hes": False,
            "requires_delivery_note": False,
            "requires_supplier_portal": False,
            "other_requirements": [],
            "risk_days_before_due": 7,
        },
    )

    class FakeSession:
        def __init__(self) -> None:
            self.results = iter(([invoice], items, [payment], [profile]))

        def scalars(self, _statement: object):
            return iter(next(self.results))

    result = build_client_portal_summary(
        FakeSession(), client=client, as_of=date(2026, 10, 25)
    )
    assert result["totals"]["invoice_count"] == 1
    assert result["totals"]["paid_amount"] == Decimal("25000.00")
    assert result["totals"]["outstanding_amount"] == Decimal("75000.00")
    assert result["totals"]["pending_client_count"] == 1
    assert result["invoices"][0]["pending_client_items"] == ["Orden de compra"]


def check_workflow_rules() -> None:
    class InvoiceStub:
        issued_on = date(2026, 10, 1)
        due_on = date(2026, 10, 31)
        status = InvoiceStatus.ISSUED

    rules = ContractRules(window_start="09:00", window_end="18:00")
    assert determine_stage(InvoiceStub(), date(2026, 10, 25), rules) == "preventive"
    saturday = datetime(2026, 10, 17, 15, 0, tzinfo=timezone.utc)
    local_schedule = next_allowed_at(saturday, rules).astimezone(ZoneInfo("America/Santiago"))
    assert local_schedule.isoformat().startswith("2026-10-19T09:00:00")


def check_complete_stage_windows() -> None:
    class InvoiceStub:
        issued_on = date(2026, 10, 1)
        due_on = date(2026, 10, 31)

    stages = [window.stage for window in stage_windows(InvoiceStub(), ContractRules())]
    assert stages == [
        "emission",
        "documental",
        "preventive",
        "due",
        "early_arrears",
        "escalation",
    ]


def main() -> None:
    checks = [
        check_rut,
        check_csv_import,
        check_xlsx_import,
        check_contact_import,
        check_sii_status_import,
        check_bank_statement_import,
        check_payability_rules,
        check_sii_deadline_alert,
        check_preventive_confirmation_contract,
        check_global_payer_profile_contract,
        check_communication_safety_contract,
        check_inbound_communication_contract,
        check_local_classification_rules,
        check_supervision_contract,
        check_openapi_contract,
        check_schema_compiles_for_postgresql,
        check_migration_chain,
        check_migrations_match_models,
        check_phase6_operational_assets,
        check_phase7_client_portal_assets,
        check_phase7_pilot_metrics_assets,
        check_phase7_pilot_operations_assets,
        check_phase8_final_report_assets,
        check_phase8_release_readiness_assets,
        check_post_plan_hardening_assets,
        check_post_plan_observability_assets,
        check_post_plan_operations_ui_assets,
        check_post_plan_improvement_plan,
        check_pm1_human_effort_assets,
        check_pilot_metrics_calculation,
        check_client_portal_calculation,
        check_workflow_rules,
        check_complete_stage_windows,
    ]
    for check in checks:
        check()
        print(f"PASS {check.__name__}")
    print(f"PASS {len(checks)} offline checks")


if __name__ == "__main__":
    main()

