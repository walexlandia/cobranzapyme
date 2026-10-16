"""Analyze a retrospective workbook without modifying it."""

import argparse
import json
from collections import defaultdict
from datetime import date, datetime
from pathlib import Path

from openpyxl import load_workbook


ROOT = Path(__file__).resolve().parents[1]
DEFAULT_WORKBOOK = ROOT / "outputs" / "recupera-datos-prueba.xlsx"


def as_date(value: object) -> date | None:
    if isinstance(value, datetime):
        return value.date()
    if isinstance(value, date):
        return value
    return None


def late_band(days: int) -> str:
    if days <= 0:
        return "on_time"
    if days <= 15:
        return "late_1_15"
    if days <= 30:
        return "late_16_30"
    return "late_31_plus"


def analyze(workbook_path: Path, *, synthetic: bool) -> dict:
    workbook = load_workbook(workbook_path, read_only=True, data_only=True)
    if "Facturas" not in workbook.sheetnames:
        raise ValueError("El libro debe contener una hoja llamada 'Facturas'")
    rows = list(workbook["Facturas"].iter_rows(min_row=2, values_only=True))
    if not rows:
        raise ValueError("La hoja 'Facturas' no contiene registros")
    if any(len(row) < 23 for row in rows):
        raise ValueError("La hoja 'Facturas' no contiene todas las columnas requeridas")
    by_client: dict[str, dict] = defaultdict(
        lambda: {
            "invoice_count": 0,
            "invoiced_amount": 0,
            "paid_count": 0,
            "on_time_count": 0,
            "dso_days": [],
            "late_distribution": defaultdict(int),
            "documentary_delay_count": 0,
            "documentary_delay_amount": 0,
        }
    )
    delayed_count = delayed_amount = 0
    documentary_count = documentary_amount = 0
    documentary_clients: set[str] = set()
    cause_summary: dict[str, dict[str, int]] = defaultdict(
        lambda: {"invoice_count": 0, "amount": 0}
    )

    for row in rows:
        client = str(row[2])
        issued_on = as_date(row[7])
        amount = int(row[9] or 0)
        paid_on = as_date(row[19])
        cause = str(row[20] or "").strip()
        days_late = int(row[21] or 0)
        paid_on_time = str(row[22] or "").strip().lower() in {"sí", "si", "yes"}
        metrics = by_client[client]
        metrics["invoice_count"] += 1
        metrics["invoiced_amount"] += amount
        if paid_on and issued_on:
            metrics["paid_count"] += 1
            metrics["dso_days"].append((paid_on - issued_on).days)
            metrics["late_distribution"][late_band(days_late)] += 1
            if paid_on_time:
                metrics["on_time_count"] += 1
        if cause:
            delayed_count += 1
            delayed_amount += amount
            cause_summary[cause]["invoice_count"] += 1
            cause_summary[cause]["amount"] += amount
            if cause.startswith("documental_"):
                documentary_count += 1
                documentary_amount += amount
                documentary_clients.add(client)
                metrics["documentary_delay_count"] += 1
                metrics["documentary_delay_amount"] += amount

    client_results = {}
    for client, metrics in sorted(by_client.items()):
        paid_count = metrics.pop("paid_count")
        on_time_count = metrics.pop("on_time_count")
        dso_days = metrics.pop("dso_days")
        distribution = dict(metrics.pop("late_distribution"))
        client_results[client] = {
            **metrics,
            "paid_invoice_count": paid_count,
            "baseline_dso": round(sum(dso_days) / len(dso_days), 2) if dso_days else None,
            "on_time_percent": round(on_time_count / paid_count * 100, 2) if paid_count else None,
            "late_distribution": distribution,
        }

    result = {
        "source": str(workbook_path.resolve()),
        "synthetic": synthetic,
        "invoice_count": len(rows),
        "delayed_invoice_count": delayed_count,
        "delayed_amount": delayed_amount,
        "documentary_delay_count": documentary_count,
        "documentary_delay_amount": documentary_amount,
        "documentary_share_by_count_percent": (
            round(documentary_count / delayed_count * 100, 2) if delayed_count else 0
        ),
        "documentary_share_by_amount_percent": (
            round(documentary_amount / delayed_amount * 100, 2) if delayed_amount else 0
        ),
        "clients_with_documentary_delays": sorted(documentary_clients),
        "cause_summary": dict(sorted(cause_summary.items())),
        "clients": client_results,
        "g1_result": (
            delayed_count > 0
            and delayed_amount > 0
            and documentary_count / delayed_count >= 0.20
            and documentary_amount / delayed_amount >= 0.20
            and len(documentary_clients) > 1
        ),
    }
    return result


def parse_args() -> argparse.Namespace:
    parser = argparse.ArgumentParser(
        description="Calcula línea base, causas de atraso y la puerta G1 sin modificar el libro."
    )
    parser.add_argument(
        "--workbook",
        type=Path,
        default=DEFAULT_WORKBOOK,
        help="Libro XLSX con la hoja Facturas (por defecto usa el dataset ficticio).",
    )
    parser.add_argument(
        "--authorized-real-data",
        action="store_true",
        help="Declara que el libro contiene datos reales cuyo análisis fue autorizado.",
    )
    return parser.parse_args()


def main() -> None:
    args = parse_args()
    if (
        args.authorized_real_data
        and args.workbook.resolve() == DEFAULT_WORKBOOK.resolve()
    ):
        raise ValueError(
            "El dataset demo es sintético y no puede declararse como información real autorizada"
        )
    result = analyze(
        args.workbook,
        synthetic=not args.authorized_real_data,
    )
    print(json.dumps(result, ensure_ascii=False, indent=2))


if __name__ == "__main__":
    main()
