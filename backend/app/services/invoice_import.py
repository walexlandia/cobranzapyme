import csv
import io
import json
import re
from datetime import date, datetime
from decimal import Decimal, InvalidOperation
from typing import Any

from openpyxl import load_workbook

from pydantic import EmailStr, TypeAdapter, ValidationError

from app.schemas import ContactImportRow, ImportIssue, InvoiceImportRow, SiiStatusImportRow
from app.utils.rut import normalize_rut


DEFAULT_ALIASES = {
    "payer_rut": ["rut pagador", "rut receptor", "rut cliente", "rut", "payer_rut"],
    "payer_name": ["razon social", "razón social", "receptor", "cliente", "payer_name"],
    "document_type": ["tipo documento", "tipo dte", "document_type"],
    "folio": ["folio", "numero", "número", "nro factura"],
    "issued_on": ["fecha emision", "fecha emisión", "issued_on"],
    "due_on": ["fecha vencimiento", "vencimiento", "due_on"],
    "total_amount": ["monto total", "total", "monto", "total_amount"],
    "currency": ["moneda", "currency"],
    "purchase_order": ["oc", "orden de compra", "purchase_order"],
}

REQUIRED_FIELDS = {"payer_rut", "payer_name", "folio", "issued_on", "due_on", "total_amount"}

CONTACT_ALIASES = {
    "payer_rut": ["rut pagador", "rut", "payer_rut"],
    "payer_name": ["razon social", "razón social", "pagador", "payer_name"],
    "name": ["nombre contacto", "contacto", "nombre", "name"],
    "role": ["rol", "cargo", "role"],
    "email": ["correo", "email", "e-mail"],
    "phone": ["telefono", "teléfono", "whatsapp", "phone"],
}

CONTACT_REQUIRED_FIELDS = {"payer_rut", "payer_name", "name", "role"}

SII_ALIASES = {
    "folio": ["folio", "numero", "número", "nro factura"],
    "document_type": ["tipo documento", "tipo dte", "document_type"],
    "sii_status": ["estado sii", "estado", "sii_status"],
    "checked_on": ["fecha consulta", "fecha estado", "checked_on"],
    "claim_deadline": ["plazo reclamo", "fecha limite reclamo", "claim_deadline"],
    "detail": ["detalle", "observacion", "observación", "detail"],
}
SII_REQUIRED_FIELDS = {"folio", "sii_status", "checked_on"}


def _normalize_header(value: Any) -> str:
    return re.sub(r"\s+", " ", str(value or "").strip().lower())


def _read_rows(filename: str, content: bytes) -> list[dict[str, Any]]:
    if filename.lower().endswith((".xlsx", ".xlsm")):
        workbook = load_workbook(io.BytesIO(content), read_only=True, data_only=True)
        sheet = workbook.active
        values = sheet.iter_rows(values_only=True)
        try:
            headers = [_normalize_header(value) for value in next(values)]
        except StopIteration:
            return []
        return [dict(zip(headers, row, strict=False)) for row in values]

    text = content.decode("utf-8-sig")
    sample = text[:4096]
    try:
        dialect = csv.Sniffer().sniff(sample, delimiters=",;\t")
    except csv.Error:
        dialect = csv.excel
    reader = csv.DictReader(io.StringIO(text), dialect=dialect)
    return [
        {_normalize_header(key): value for key, value in row.items() if key is not None}
        for row in reader
    ]


def _resolve_mapping(headers: set[str], raw_mapping: str | None) -> dict[str, str]:
    if raw_mapping:
        supplied = json.loads(raw_mapping)
        return {field: _normalize_header(header) for field, header in supplied.items()}
    mapping: dict[str, str] = {}
    for field, aliases in DEFAULT_ALIASES.items():
        match = next((alias for alias in aliases if _normalize_header(alias) in headers), None)
        if match:
            mapping[field] = _normalize_header(match)
    return mapping


def _parse_date(value: Any) -> date:
    if isinstance(value, datetime):
        return value.date()
    if isinstance(value, date):
        return value
    raw = str(value or "").strip()
    for pattern in ("%Y-%m-%d", "%d-%m-%Y", "%d/%m/%Y", "%Y/%m/%d"):
        try:
            return datetime.strptime(raw, pattern).date()
        except ValueError:
            continue
    raise ValueError("fecha inválida; use AAAA-MM-DD o DD-MM-AAAA")


def _parse_amount(value: Any) -> Decimal:
    if isinstance(value, (int, float, Decimal)):
        amount = Decimal(str(value))
    else:
        raw = re.sub(r"[^0-9,.-]", "", str(value or ""))
        if "," in raw and "." in raw:
            raw = raw.replace(".", "").replace(",", ".")
        elif "," in raw:
            raw = raw.replace(",", ".")
        elif raw.count(".") > 1:
            raw = raw.replace(".", "")
        elif raw.count(".") == 1 and len(raw.rsplit(".", 1)[1]) == 3:
            raw = raw.replace(".", "")
        try:
            amount = Decimal(raw)
        except InvalidOperation as exc:
            raise ValueError("monto inválido") from exc
    if amount <= 0:
        raise ValueError("el monto debe ser mayor a cero")
    return amount.quantize(Decimal("0.01"))


def parse_invoice_file(
    filename: str,
    content: bytes,
    raw_mapping: str | None = None,
) -> tuple[list[InvoiceImportRow], list[ImportIssue]]:
    raw_rows = _read_rows(filename, content)
    if not raw_rows:
        return [], [ImportIssue(row=1, message="El archivo no contiene filas")]
    mapping = _resolve_mapping(set(raw_rows[0]), raw_mapping)
    missing = REQUIRED_FIELDS - set(mapping)
    if missing:
        return [], [
            ImportIssue(row=1, field=field, message=f"Falta mapear la columna {field}")
            for field in sorted(missing)
        ]

    valid_rows: list[InvoiceImportRow] = []
    issues: list[ImportIssue] = []
    seen: set[tuple[str, str]] = set()
    for index, raw in enumerate(raw_rows, start=2):
        if not any(value not in (None, "") for value in raw.values()):
            continue
        try:
            document_type = str(raw.get(mapping.get("document_type", "")) or "factura").strip().lower()
            folio = str(raw.get(mapping["folio"]) or "").strip()
            if not folio:
                raise ValueError("folio vacío")
            key = (document_type, folio)
            if key in seen:
                raise ValueError("factura duplicada dentro del archivo")
            issued_on = _parse_date(raw.get(mapping["issued_on"]))
            due_on = _parse_date(raw.get(mapping["due_on"]))
            if due_on < issued_on:
                raise ValueError("el vencimiento es anterior a la emisión")
            row = InvoiceImportRow(
                row=index,
                payer_rut=normalize_rut(str(raw.get(mapping["payer_rut"]) or "")),
                payer_name=str(raw.get(mapping["payer_name"]) or "").strip(),
                document_type=document_type,
                folio=folio,
                issued_on=issued_on,
                due_on=due_on,
                total_amount=_parse_amount(raw.get(mapping["total_amount"])),
                currency=str(raw.get(mapping.get("currency", "")) or "CLP").strip().upper(),
                purchase_order=(
                    str(raw.get(mapping.get("purchase_order", ""))).strip()
                    if raw.get(mapping.get("purchase_order", "")) not in (None, "")
                    else None
                ),
            )
            if not row.payer_name:
                raise ValueError("razón social vacía")
            seen.add(key)
            valid_rows.append(row)
        except (ValueError, TypeError) as exc:
            issues.append(ImportIssue(row=index, message=str(exc)))
    return valid_rows, issues


def parse_contact_file(
    filename: str,
    content: bytes,
    raw_mapping: str | None = None,
) -> tuple[list[ContactImportRow], list[ImportIssue]]:
    raw_rows = _read_rows(filename, content)
    if not raw_rows:
        return [], [ImportIssue(row=1, message="El archivo no contiene filas")]
    if raw_mapping:
        mapping = _resolve_mapping(set(raw_rows[0]), raw_mapping)
    else:
        mapping = {}
        headers = set(raw_rows[0])
        for field, aliases in CONTACT_ALIASES.items():
            match = next((alias for alias in aliases if _normalize_header(alias) in headers), None)
            if match:
                mapping[field] = _normalize_header(match)
    missing = CONTACT_REQUIRED_FIELDS - set(mapping)
    if missing:
        return [], [
            ImportIssue(row=1, field=field, message=f"Falta mapear la columna {field}")
            for field in sorted(missing)
        ]

    role_map = {
        "cuentas por pagar": "accounts_payable",
        "cxp": "accounts_payable",
        "accounts payable": "accounts_payable",
        "solicitante": "requester",
        "administrador de contrato": "requester",
        "requester": "requester",
        "otro": "other",
        "other": "other",
    }
    email_adapter = TypeAdapter(EmailStr)
    valid_rows: list[ContactImportRow] = []
    issues: list[ImportIssue] = []
    seen: set[tuple[str, str, str]] = set()
    for index, raw in enumerate(raw_rows, start=2):
        if not any(value not in (None, "") for value in raw.values()):
            continue
        try:
            payer_rut = normalize_rut(str(raw.get(mapping["payer_rut"]) or ""))
            payer_name = str(raw.get(mapping["payer_name"]) or "").strip()
            name = str(raw.get(mapping["name"]) or "").strip()
            raw_role = _normalize_header(raw.get(mapping["role"]))
            role = role_map.get(raw_role)
            if not payer_name or not name:
                raise ValueError("pagador y nombre del contacto son obligatorios")
            if role is None:
                raise ValueError("rol inválido; use cuentas por pagar, solicitante u otro")
            raw_email = str(raw.get(mapping.get("email", "")) or "").strip().lower()
            email = str(email_adapter.validate_python(raw_email)) if raw_email else None
            phone = str(raw.get(mapping.get("phone", "")) or "").strip() or None
            if not email and not phone:
                raise ValueError("el contacto necesita correo o teléfono")
            key = (payer_rut, email or "", phone or "")
            if key in seen:
                raise ValueError("contacto duplicado dentro del archivo")
            seen.add(key)
            valid_rows.append(
                ContactImportRow(
                    row=index,
                    payer_rut=payer_rut,
                    payer_name=payer_name,
                    name=name,
                    role=role,
                    email=email,
                    phone=phone,
                )
            )
        except (ValueError, TypeError, ValidationError) as exc:
            issues.append(ImportIssue(row=index, message=str(exc)))
    return valid_rows, issues


def parse_sii_status_file(
    filename: str,
    content: bytes,
    raw_mapping: str | None = None,
) -> tuple[list[SiiStatusImportRow], list[ImportIssue]]:
    raw_rows = _read_rows(filename, content)
    if not raw_rows:
        return [], [ImportIssue(row=1, message="El archivo no contiene filas")]
    if raw_mapping:
        supplied = json.loads(raw_mapping)
        mapping = {field: _normalize_header(header) for field, header in supplied.items()}
    else:
        mapping = {}
        headers = set(raw_rows[0])
        for field, aliases in SII_ALIASES.items():
            match = next((alias for alias in aliases if _normalize_header(alias) in headers), None)
            if match:
                mapping[field] = _normalize_header(match)
    missing = SII_REQUIRED_FIELDS - set(mapping)
    if missing:
        return [], [
            ImportIssue(row=1, field=field, message=f"Falta mapear la columna {field}")
            for field in sorted(missing)
        ]

    status_map = {
        "aceptada": "accepted",
        "aceptado": "accepted",
        "accepted": "accepted",
        "reclamada": "claimed",
        "reclamado": "claimed",
        "claimed": "claimed",
        "pendiente": "pending",
        "pending": "pending",
        "rechazada": "rejected",
        "rechazado": "rejected",
        "rejected": "rejected",
    }
    valid_rows: list[SiiStatusImportRow] = []
    issues: list[ImportIssue] = []
    seen: set[tuple[str, str]] = set()
    for index, raw in enumerate(raw_rows, start=2):
        if not any(value not in (None, "") for value in raw.values()):
            continue
        try:
            folio = str(raw.get(mapping["folio"]) or "").strip()
            if not folio:
                raise ValueError("folio vacío")
            document_type = str(
                raw.get(mapping.get("document_type", "")) or "factura"
            ).strip().lower()
            key = (document_type, folio)
            if key in seen:
                raise ValueError("factura duplicada dentro del archivo")
            raw_status = _normalize_header(raw.get(mapping["sii_status"]))
            sii_status = status_map.get(raw_status)
            if sii_status is None:
                raise ValueError("estado SII inválido")
            claim_deadline_raw = raw.get(mapping.get("claim_deadline", ""))
            valid_rows.append(
                SiiStatusImportRow(
                    row=index,
                    folio=folio,
                    document_type=document_type,
                    sii_status=sii_status,
                    checked_on=_parse_date(raw.get(mapping["checked_on"])),
                    claim_deadline=(
                        _parse_date(claim_deadline_raw)
                        if claim_deadline_raw not in (None, "")
                        else None
                    ),
                    detail=(
                        str(raw.get(mapping.get("detail", ""))).strip() or None
                        if raw.get(mapping.get("detail", "")) not in (None, "")
                        else None
                    ),
                )
            )
            seen.add(key)
        except (ValueError, TypeError) as exc:
            issues.append(ImportIssue(row=index, message=str(exc)))
    return valid_rows, issues

